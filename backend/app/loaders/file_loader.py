"""Loads various file types into LangChain Documents."""
from pathlib import Path
from langchain_core.documents import Document as LCDocument
from langchain_community.document_loaders import PyPDFLoader, TextLoader, CSVLoader

SUPPORTED_EXTENSIONS = {
    '.pdf', '.md', '.markdown', '.csv', '.html', '.htm', '.xlsx', '.ods',
}


def load_file(path: Path) -> tuple[list[LCDocument], str]:
    """Returns (docs, source_type). source_type is the canonical format name."""
    ext = path.suffix.lower()

    if ext == '.pdf':
        return PyPDFLoader(str(path)).load(), 'pdf'

    if ext == '.xlsx':
        return _load_xlsx(path), 'xlsx'

    if ext == '.ods':
        return _load_ods(path), 'xlsx'

    if ext == '.csv':
        return CSVLoader(str(path)).load(), 'csv'

    if ext in ('.html', '.htm'):
        return _load_html_file(path), 'html'

    if ext in ('.md', '.markdown'):
        docs = TextLoader(str(path), encoding='utf-8').load()
        # Assign page numbers so downstream metadata works
        for i, doc in enumerate(docs):
            doc.metadata.setdefault('page', i)
        return docs, 'md'

    raise ValueError(f"Unsupported extension: {ext}")


def _load_xlsx(path: Path) -> list[LCDocument]:
    """Treats the first non-empty row of each sheet as the column header,
    kept separate so it can be repeated on every downstream chunk."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    docs = []
    for idx, sheet in enumerate(wb.worksheets, 1):
        raw_rows = [
            row for row in sheet.iter_rows(values_only=True)
            if any(c is not None for c in row)
        ]
        if not raw_rows:
            continue
        header, *data_rows = raw_rows
        header_line = ', '.join(str(c) for c in header if c is not None)
        data_lines = [', '.join(str(c) for c in row if c is not None) for row in data_rows]
        docs.append(LCDocument(
            page_content=f'Sheet: {sheet.title}\nColumns: {header_line}\n\n' + '\n'.join(data_lines),
            metadata={'page': idx, 'source': str(path)},
        ))
    return docs or [LCDocument(page_content='Empty spreadsheet', metadata={'page': 1})]


def _load_ods(path: Path) -> list[LCDocument]:
    """Same output shape as _load_xlsx (Sheet:/Columns: header + one line per
    row) so it reuses the existing xlsx chunking strategy unchanged."""
    from odf.opendocument import load as load_ods
    from odf.table import Table, TableRow, TableCell
    from odf import teletype

    doc = load_ods(str(path))
    docs = []
    for idx, sheet in enumerate(doc.spreadsheet.getElementsByType(Table), 1):
        raw_rows = []
        for row in sheet.getElementsByType(TableRow):
            values = [teletype.extractText(cell) for cell in row.getElementsByType(TableCell)]
            if any(v.strip() for v in values):
                raw_rows.append(values)
        if not raw_rows:
            continue
        header, *data_rows = raw_rows
        header_line = ', '.join(v for v in header if v)
        data_lines = [', '.join(v for v in row if v) for row in data_rows]
        docs.append(LCDocument(
            page_content=f'Sheet: {sheet.getAttribute("name")}\nColumns: {header_line}\n\n' + '\n'.join(data_lines),
            metadata={'page': idx, 'source': str(path)},
        ))
    return docs or [LCDocument(page_content='Empty spreadsheet', metadata={'page': 1})]


def _load_html_file(path: Path) -> list[LCDocument]:
    import trafilatura
    content = path.read_text(encoding='utf-8', errors='ignore')
    text = trafilatura.extract(content, output_format='markdown', include_formatting=True)

    if not text:
        # Fallback for fragments trafilatura can't confidently extract
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(content, 'lxml')
        for tag in soup(['script', 'style', 'nav', 'footer']):
            tag.decompose()
        text = soup.get_text(separator='\n', strip=True)

    return [LCDocument(page_content=text, metadata={'page': 1, 'source': str(path)})]